"""Yevmiye ve Kebir XML'lerinden mali yıl başına Excel (Yevmiye, Kebir, Mizan) üretir."""
from __future__ import annotations

import datetime as dt
import xml.etree.ElementTree as ET
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from edefter import Doc

COR = "{http://www.xbrl.org/int/gl/cor/2006-10-25}"
BUS = "{http://www.xbrl.org/int/gl/bus/2006-10-25}"

# GİB yevmiye.xslt'deki karşılıklar
DOC_TYPES = {"check": "Çek", "invoice": "Fatura", "order-customer": "Müşteri Siparişi",
             "order-vendor": "Satıcı Siparişi", "voucher": "Senet", "shipment": "Navlun", "receipt": "Makbuz"}

MONEY = "#,##0.00"
DATE = "DD.MM.YYYY"
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F5D7C")
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="E8EEF3")


def _t(el, path: str) -> str:
    v = el.find(path)
    return (v.text or "").strip() if v is not None and v.text else ""


def _date(s: str):
    try:
        return dt.date.fromisoformat(s[:10])
    except ValueError:
        return s or None


def parse_lines(doc: Doc) -> list[dict]:
    """XBRL-GL entryDetail satırlarını düz sözlüklere çevirir (başlık alanları satıra taşınır)."""
    root = ET.parse(doc.src).getroot()
    rows = []
    for header in root.iter(COR + "entryHeader"):
        h = {
            "yev_no": _t(header, COR + "entryNumberCounter"),
            "yev_tarih": _date(_t(header, COR + "enteredDate")),
            "fis_no": _t(header, COR + "entryNumber"),
            "fis_aciklama": _t(header, COR + "entryComment"),
        }
        for d in header.iter(COR + "entryDetail"):
            amount = Decimal(_t(d, COR + "amount") or "0")
            debit = _t(d, COR + "debitCreditCode").upper().startswith("D")
            dtype = _t(d, COR + "documentType")
            rows.append({
                **h,
                "ay": f"{doc.year}-{doc.month}",
                "satir_no": _t(d, COR + "lineNumber"),
                # Kebir satırında yevmiye no / fiş no satırın kendisinde durur
                "yev_no_satir": _t(d, COR + "lineNumberCounter"),
                "hesap": _t(d, f"{COR}account/{COR}accountMainID"),
                "hesap_adi": _t(d, f"{COR}account/{COR}accountMainDescription"),
                "alt_hesap": _t(d, f"{COR}account/{COR}accountSub/{COR}accountSubID"),
                "alt_hesap_adi": _t(d, f"{COR}account/{COR}accountSub/{COR}accountSubDescription"),
                "aciklama": _t(d, COR + "detailComment"),
                "tarih": _date(_t(d, COR + "postingDate")),
                "belge_turu": DOC_TYPES.get(dtype, _t(d, COR + "documentTypeDescription") or dtype),
                "belge_no": _t(d, COR + "documentNumber"),
                "belge_tarihi": _date(_t(d, COR + "documentDate")),
                "fis_ref": _t(d, COR + "documentReference"),
                "odeme": _t(d, BUS + "paymentMethod"),
                "borc": amount if debit else Decimal(0),
                "alacak": Decimal(0) if debit else amount,
            })
    return rows


# ---------------------------------------------------------------- sayfa yazımı

def _sheet(wb: Workbook, title: str, headers: list[tuple[str, int, str | None]]):
    ws = wb.create_sheet(title)
    ws.append([h for h, _, _ in headers])
    for i, (_, width, _) in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    return ws


def _finish(ws, headers, total_row: list | None = None):
    fmts = [f for _, _, f in headers]
    for row in ws.iter_rows(min_row=2):
        for c, f in zip(row, fmts):
            if f:
                c.number_format = f
    ws.auto_filter.ref = ws.dimensions
    if total_row:
        ws.append(total_row)
        for c, f in zip(ws[ws.max_row], fmts):
            c.font, c.fill = TOTAL_FONT, TOTAL_FILL
            if f:
                c.number_format = f


def _num(x: Decimal) -> float:
    return float(x)


def _balance(b: Decimal, a: Decimal):
    diff = b - a
    return (_num(diff), 0.0) if diff >= 0 else (0.0, _num(-diff))


def write_year(target: Path, yevmiye: list[dict], kebir: list[dict], title: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    # Yevmiye
    yh = [("Ay", 9, None), ("Yevmiye No", 10, None), ("Yevmiye Tarihi", 12, DATE), ("Fiş No", 12, None),
          ("Fiş Açıklaması", 30, None), ("Hesap Kodu", 10, None), ("Hesap Adı", 26, None),
          ("Alt Hesap Kodu", 15, None), ("Alt Hesap Adı", 30, None), ("Açıklama", 40, None),
          ("Belge Türü", 11, None), ("Belge No", 18, None), ("Belge Tarihi", 12, DATE),
          ("Ödeme Şekli", 12, None), ("Borç", 16, MONEY), ("Alacak", 16, MONEY)]
    ws = _sheet(wb, "Yevmiye", yh)
    tb = ta = Decimal(0)
    for r in yevmiye:
        ws.append([r["ay"], _int(r["yev_no"]), r["yev_tarih"], r["fis_no"], r["fis_aciklama"], r["hesap"],
                   r["hesap_adi"], r["alt_hesap"], r["alt_hesap_adi"], r["aciklama"], r["belge_turu"],
                   r["belge_no"], r["belge_tarihi"], r["odeme"], _num(r["borc"]), _num(r["alacak"])])
        tb += r["borc"]
        ta += r["alacak"]
    _finish(ws, yh, ["TOPLAM"] + [None] * 13 + [_num(tb), _num(ta)])

    # Kebir: hesap → alt hesap → ay sırası; yıl içinde kümülatif
    kh = [("Hesap Kodu", 10, None), ("Hesap Adı", 26, None), ("Alt Hesap Kodu", 15, None),
          ("Alt Hesap Adı", 30, None), ("Tarih", 12, DATE), ("Yevmiye No", 10, None), ("Fiş No", 12, None),
          ("Açıklama", 40, None), ("Belge Türü", 11, None), ("Belge No", 18, None), ("Borç", 16, MONEY),
          ("Alacak", 16, MONEY), ("Toplam Borç", 16, MONEY), ("Toplam Alacak", 16, MONEY),
          ("Borç Bakiye", 16, MONEY), ("Alacak Bakiye", 16, MONEY)]
    ws = _sheet(wb, "Kebir", kh)
    ordered = sorted(enumerate(kebir), key=lambda x: (_acc_key(x[1]["hesap"]), _acc_key(x[1]["alt_hesap"]), x[0]))
    run: dict[str, list[Decimal]] = defaultdict(lambda: [Decimal(0), Decimal(0)])
    for _, r in ordered:
        acc = r["alt_hesap"] or r["hesap"]
        run[acc][0] += r["borc"]
        run[acc][1] += r["alacak"]
        b, a = run[acc]
        ws.append([r["hesap"], r["hesap_adi"], r["alt_hesap"], r["alt_hesap_adi"], r["tarih"],
                   _int(r["yev_no_satir"]), r["fis_ref"], r["aciklama"], r["belge_turu"], r["belge_no"],
                   _num(r["borc"]), _num(r["alacak"]), _num(b), _num(a), *_balance(b, a)])
    kb = sum((r["borc"] for r in kebir), Decimal(0))
    ka = sum((r["alacak"] for r in kebir), Decimal(0))
    _finish(ws, kh, ["TOPLAM"] + [None] * 9 + [_num(kb), _num(ka), None, None, None, None])

    # Mizan (Kebir'den; Kebir yoksa Yevmiye'den)
    source = kebir or yevmiye
    for title_, key in (("Mizan (Ana Hesap)", "hesap"), ("Mizan (Alt Hesap)", "alt_hesap")):
        mh = [("Hesap Kodu", 15, None), ("Hesap Adı", 40, None), ("Borç", 18, MONEY), ("Alacak", 18, MONEY),
              ("Borç Bakiye", 18, MONEY), ("Alacak Bakiye", 18, MONEY)]
        ws = _sheet(wb, title_, mh)
        acc: dict[str, list] = {}
        for r in source:
            code = r[key] or r["hesap"]
            name = r[key + "_adi"] if r[key] else r["hesap_adi"]
            e = acc.setdefault(code, [name, Decimal(0), Decimal(0)])
            e[1] += r["borc"]
            e[2] += r["alacak"]
        sb = sa = bb = ba = Decimal(0)
        for code in sorted(acc, key=_acc_key):
            name, b, a = acc[code]
            ws.append([code, name, _num(b), _num(a), *_balance(b, a)])
            sb, sa = sb + b, sa + a
            bb, ba = (bb + b - a, ba) if b >= a else (bb, ba + a - b)
        _finish(ws, mh, ["TOPLAM", None, _num(sb), _num(sa), _num(bb), _num(ba)])

    wb.properties.title = title
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)


def _int(s: str):
    return int(s) if s.isdigit() else s


def _acc_key(code: str):
    # "101.00.02" gibi kodları parça parça sayısal sırala; harfli parçalar sayılardan sonra
    return [(0, int(p), "") if p.isdigit() else (1, 0, p) for p in code.split(".")] if code else []


def export(out_root: Path, docs: list[Doc]) -> list[Path]:
    """VKN + mali yıl başına '<yıl> Defter.xlsx' yazar. İptal edilenler dahil edilmez."""
    groups: dict[tuple[str, str], dict[str, list[Doc]]] = defaultdict(lambda: {"Y": [], "K": []})
    for d in docs:
        if d.kind in ("Y", "K") and not d.gib and not d.cancelled:
            groups[(d.vkn, d.year)][d.kind].append(d)
    made = []
    for (vkn, year), kinds in sorted(groups.items()):
        ykey = lambda d: (d.month, d.part)
        yev = [r for d in sorted(kinds["Y"], key=ykey) for r in parse_lines(d)]
        keb = [r for d in sorted(kinds["K"], key=ykey) for r in parse_lines(d)]
        target = out_root / vkn / year / f"{year} Defter (Yevmiye-Kebir-Mizan).xlsx"
        write_year(target, yev, keb, f"{vkn} {year} e-Defter")
        made.append(target)
    return made
